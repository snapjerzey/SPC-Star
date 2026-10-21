"""Regenerate repository SOP references from the in-app SOP sections."""
from html.parser import HTMLParser
from pathlib import Path
import argparse
import json
from shutil import copy2


ROOT = Path(__file__).resolve().parents[1]
SECTIONS = {
    'sopEnglishBody': ('operator-en.md', 'Operator SOP'),
    'sopSpanishBody': ('operator-es.md', 'SOP del operador'),
    'lineTechSopEnglishBody': ('line-tech-en.md', 'Line Tech SOP'),
    'lineTechSopSpanishBody': ('line-tech-es.md', 'SOP del técnico de línea'),
    'qaSopEnglishBody': ('qa-en.md', 'QA SOP'),
    'qaSopSpanishBody': ('qa-es.md', 'SOP de QA'),
}


class SopParser(HTMLParser):
    def __init__(self, source=''):
        super().__init__(convert_charrefs=True)
        self.sections = {}
        self.active = None
        self.depth = 0
        self.block = None
        self.text = []
        self.number = 0
        self.bounds = {}
        self.line_offsets = [0]
        for line in source.splitlines(keepends=True):
            self.line_offsets.append(self.line_offsets[-1] + len(line))

    def source_position(self):
        line, col = self.getpos()
        return self.line_offsets[line - 1] + col

    def handle_starttag(self, tag, attrs):
        if tag == 'div':
            section_id = dict(attrs).get('id')
            if section_id in SECTIONS:
                assert self.active is None and section_id not in self.sections
                self.active = section_id
                self.sections[section_id] = []
                self.depth = 1
                self.bounds[section_id] = [self.source_position() + len(self.get_starttag_text()), None]
            elif self.active:
                self.depth += 1
        if not self.active:
            return
        if tag == 'ol':
            self.number = 0
        if tag in ('h3', 'p', 'li'):
            assert self.block is None, f'Nested SOP text block: {tag}'
            self.block = tag
            self.text = []

    def handle_data(self, data):
        if self.active and self.block:
            self.text.append(data)

    def handle_endtag(self, tag):
        if not self.active:
            return
        if tag == self.block:
            text = ' '.join(''.join(self.text).split())
            if tag == 'h3':
                line = '## ' + text
            elif tag == 'li':
                self.number += 1
                line = f'{self.number}. {text}'
            else:
                line = text
            self.sections[self.active].append((tag, line))
            self.block = None
        if tag == 'div':
            self.depth -= 1
            if self.depth == 0:
                self.bounds[self.active][1] = self.source_position()
                self.active = None


def publish(content_path):
    path = ROOT / 'src/SPCStar.Api/wwwroot/index.html'
    source = path.read_text(encoding='utf-8')
    parser = SopParser(source)
    parser.feed(source)
    content = json.loads(Path(content_path).read_text(encoding='utf-8'))
    assert set(content) == set(SECTIONS) == set(parser.bounds)
    for section_id, (start, end) in sorted(parser.bounds.items(), key=lambda item: item[1][0], reverse=True):
        assert end is not None
        source = source[:start] + '\n' + content[section_id] + source[end:]
    path.write_text(source, encoding='utf-8')
    print('Published six in-app SOP sections')


def documents(downloads, suffix=''):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source = (ROOT / 'src/SPCStar.Api/wwwroot/index.html').read_text(encoding='utf-8')
    parser = SopParser(source)
    parser.feed(source)
    output = ROOT / 'docs/sops/word'
    output.mkdir(exist_ok=True)
    for section_id, (filename, _) in SECTIONS.items():
        role = {'operator': 'Operator', 'line-tech': 'Line Tech', 'qa': 'QA'}[filename[:-6]]
        language = 'English' if filename.endswith('-en.md') else 'Spanish'
        title = f'SPC-Star {role} SOP'
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = section.bottom_margin = Inches(.7)
        section.left_margin = section.right_margin = Inches(.8)
        for name, size in [('Normal', 10.5), ('Title', 22), ('Heading 1', 12)]:
            style = document.styles[name]
            style.font.name = 'Calibri'
            style.font.size = Pt(size)
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.line_spacing = 1.05
            for border in list(style.element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
        document.styles['Heading 1'].paragraph_format.space_before = Pt(10)
        if filename == 'operator-es.md':
            document.styles['Normal'].paragraph_format.space_after = Pt(4)
        document.add_paragraph(title + ' ' + language, 'Title')
        document.core_properties.title = title + ' ' + language
        document.core_properties.author = ''
        for tag, line in parser.sections[section_id]:
            if tag == 'h3':
                document.add_paragraph(line.removeprefix('## '), 'Heading 1')
            else:
                paragraph = document.add_paragraph(line)
                if tag == 'li':
                    paragraph.paragraph_format.left_indent = Inches(.18)
                    paragraph.paragraph_format.first_line_indent = Inches(-.18)
                    paragraph.paragraph_format.keep_together = True
        footer = section.footer.paragraphs[0]
        footer.add_run(title + ' | ' + language + ' | ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
        for run in footer.runs:
            run.font.size = Pt(8)
        target = output / f'{title} - {language}{suffix}.docx'
        document.save(target)
        copy2(target, Path(downloads) / target.name)
        print(target)


def export():
    source = (ROOT / 'src/SPCStar.Api/wwwroot/index.html').read_text(encoding='utf-8')
    parser = SopParser(source)
    parser.feed(source)
    assert set(parser.sections) == set(SECTIONS), 'Missing role/language SOP'
    output_dir = ROOT / 'docs/sops'
    output_dir.mkdir(exist_ok=True)
    for section_id, (filename, title) in SECTIONS.items():
        blocks = parser.sections[section_id]
        assert sum(tag == 'h3' for tag, _ in blocks) >= 5, f'Incomplete SOP: {section_id}'
        content = '# SPC-Star ' + title + '\n\n'
        previous = None
        for tag, line in blocks:
            if previous is not None and not (previous == 'li' and tag == 'li'):
                content += '\n'
            content += line + '\n'
            previous = tag
        (output_dir / filename).write_text(content, encoding='utf-8')
        print(filename)


if __name__ == '__main__':
    cli = argparse.ArgumentParser()
    cli.add_argument('--publish-json')
    cli.add_argument('--downloads')
    cli.add_argument('--document-suffix', default='')
    args = cli.parse_args()
    if args.publish_json:
        publish(args.publish_json)
    if args.downloads:
        documents(args.downloads, args.document_suffix)
    if not args.publish_json and not args.downloads:
        export()
